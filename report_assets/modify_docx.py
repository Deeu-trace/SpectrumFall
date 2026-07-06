"""
Modify docx directly (without going through md) to highlight 4 innovation features:
- I-01 生存模式 (Survival Mode)
- I-02 主题系统 (Theme System)
- I-04 谱面密度归一化 (Chart Density Normalization)
- I-05 难度自动评级 (Auto Difficulty Rating)

Strategy:
1. Insert a new subsection "1.2.4 创新功能亮点详述" after the innovation features table
2. Modify the abstract to emphasize innovation features
"""
import sys
sys.path.insert(0, r"E:\Lib\site-packages")

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from copy import deepcopy

SRC = r"e:\SpectrumFall\课程设计报告.docx"
DST = r"e:\SpectrumFall\课程设计报告.docx"

doc = Document(SRC)
body = doc.element.body

# Collect elements in order
elements = []
for child in body.iterchildren():
    elements.append(child)

# ===== Step 1: Locate the innovation features table and the "1.3 非功能需求" paragraph =====
innovation_table_idx = None
next_section_idx = None  # The paragraph "1.3 非功能需求" - we insert before it

for i, el in enumerate(elements):
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        from docx.text.paragraph import Paragraph
        p = Paragraph(el, doc)
        text = p.text.strip()
        if text == '1.3 非功能需求':
            next_section_idx = i
            break

if next_section_idx is None:
    print("ERROR: Could not find '1.3 非功能需求'")
    sys.exit(1)

# Find the paragraph "在任务书要求之外" then the first table after it
anchor_para = None
anchor_table = None
anchor_para_idx = None
for i, el in enumerate(elements):
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        from docx.text.paragraph import Paragraph
        p = Paragraph(el, doc)
        if '在任务书要求之外' in p.text:
            anchor_para = el
            anchor_para_idx = i
            break

if anchor_para_idx is not None:
    # search forward for the first table
    for j in range(anchor_para_idx + 1, len(elements)):
        el = elements[j]
        tag = el.tag.split('}')[-1]
        if tag == 'tbl':
            anchor_table = el
            break
        elif tag == 'p':
            # if we hit another paragraph that looks like a section heading, stop
            from docx.text.paragraph import Paragraph
            p = Paragraph(el, doc)
            t = p.text.strip()
            if t.startswith('1.3') or t.startswith('1.2.4'):
                break

if anchor_para is None or anchor_table is None:
    print("ERROR: Could not locate innovation table")
    print(f"  anchor_para found: {anchor_para is not None}")
    print(f"  anchor_table found: {anchor_table is not None}")
    sys.exit(1)

print(f"Found innovation table and anchor paragraph")

# ===== Step 2: Build new content to insert after the table =====
# We will insert paragraphs after the table element (before "1.3 非功能需求")
# Using python-docx, we add paragraphs to the end then move them, OR we use XML insertion

def make_heading(doc, text, level=3):
    """Create a heading paragraph element"""
    p = doc.add_paragraph()
    p.style = doc.styles[f'Heading {level}']
    run = p.add_run(text)
    return p._element

def make_para(doc, text, bold=False, italic=False):
    """Create a body text paragraph"""
    p = doc.add_paragraph()
    p.style = doc.styles['Body Text']
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p._element

def make_para_with_runs(doc, runs_spec):
    """Create a paragraph with multiple runs (list of (text, bold, italic) tuples)"""
    p = doc.add_paragraph()
    p.style = doc.styles['Body Text']
    for text, bold, italic in runs_spec:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
    return p._element

# Build all new elements
new_elements = []

# 1.2.4 subsection heading
new_elements.append(make_heading(doc, "1.2.4 创新功能亮点详述", level=3))

# Intro
new_elements.append(make_para(doc,
    "本项目在 8 项创新功能中，以下 4 项为最具代表性、技术含量最高的核心创新，"
    "体现了项目在游戏性、用户体验、算法设计上的深入思考与工程实现。"
))

# (1) 生存模式
new_elements.append(make_para_with_runs(doc, [
    ("（1）生存模式（I-01）—— 游戏性深化", True, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("设计思路：", True, False),
    ("传统节奏游戏缺乏失败惩罚，玩家压力感不足。本系统引入血量条机制，将节奏游戏从「刷分」模式升级为「生存」模式，显著提升游戏紧张感。", False, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("技术实现：", True, False),
    ("血量初始值与最大值均为 100；每次 Miss 扣除 10 点血量；连击达到 50/100/200 里程碑时分别回血 5/10/15 点（防止无脑刷连击）；血量归零立即 Game Over，结算页显示「生存失败」；血条颜色随血量动态变化：绿（>60%）→ 黄（30-60%）→ 红（<30%），直观反映危险状态。", False, False),
]))

# (2) 主题系统
new_elements.append(make_para_with_runs(doc, [
    ("（2）主题系统（I-02）—— 视觉个性化", True, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("设计思路：", True, False),
    ("不同曲风的音频适合不同的视觉风格。本系统实现「情绪自动分类 + 主题匹配」机制，让画面自动适配音乐氛围，无需用户手动调色。", False, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("技术实现：", True, False),
    ("5 种预设主题（激昂/欢快/舒缓/低沉/通用）分别对应红橙、黄绿、青蓝、紫黑、灰色调色板；情绪分类算法基于 BPM（慢/中/快）、低频能量比（低音强弱）、整体能量（动态范围）三维特征加权决策；双轨道调色板设计——菜单面板用 QSS 渲染，游戏绘制用运行时色值，同一主题两套色值协同；ThemeManager 采用单例模式，通过信号槽全局通知主题变更；自定义调色器支持 9 色块实时预览与 JSON 导入/导出。", False, False),
]))

# (3) 谱面密度归一化
new_elements.append(make_para_with_runs(doc, [
    ("（3）谱面密度归一化（I-04）—— 算法创新", True, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("设计思路：", True, False),
    ("不同歌曲时长差异巨大（30 秒到 10 分钟），若按固定密度生成音符，长歌过密、短歌过稀。需要根据时长自适应调节，保证不同歌曲的游戏体验一致。", False, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("技术实现：", True, False),
    ("目标密度设定为 100 音符/分钟（经验值，兼顾挑战性和可玩性）；算法流程：①计算目标音符数 = 时长(秒) × 100 / 60；②若实际音符数 > 目标数，按均匀间隔剔除多余 TAP（保留 HOLD 不变）；③若实际音符数 < 目标数，在节拍点中插入补充音符（优先选择强拍）；剔除策略采用等间隔剔除，避免连续剔除导致谱面断层。", False, False),
]))

# (4) 难度自动评级
new_elements.append(make_para_with_runs(doc, [
    ("（4）难度自动评级（I-05）—— 数据驱动设计", True, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("设计思路：", True, False),
    ("传统节奏游戏需要人工标注难度，效率低且主观。本系统实现全自动难度评估，基于谱面特征数据驱动计算难度星级。", False, False),
]))
new_elements.append(make_para_with_runs(doc, [
    ("技术实现：", True, False),
    ("四维特征提取：①音符密度（音符数/分钟）；②短间隔占比（< 200ms 的连续音符比例，反映高难度连打）；③BPM（音乐速度）；④最长密集段长度（连续高密度音符的秒数）。加权评分公式：difficulty = 0.3×密度 + 0.25×短间隔 + 0.2×BPM + 0.25×密集段。归一化到 1-10 星：1-2 星（入门）、3-4 星（简单）、5-6 星（中等）、7-8 星（困难）、9-10 星（极限）。玩家选歌前能直观了解难度，避免选到超出能力的歌曲。", False, False),
]))

# Summary sentence
new_elements.append(make_para_with_runs(doc, [
    ("以上 4 项创新功能在游戏性、视觉个性化、算法设计、数据驱动四个维度形成了项目的核心亮点，", False, False),
    ("在任务书要求之外展现了较强的工程创新能力和综合设计能力", True, False),
    ("。", False, False),
]))

# ===== Step 3: Insert new elements after the innovation table =====
# The table element is anchor_table. We insert new_elements right after it.
# In XML, we use addnext() to insert sibling after an element
current = anchor_table
for new_el in new_elements:
    current.addnext(new_el)
    current = new_el  # chain: each new element goes after the previous

print(f"Inserted {len(new_elements)} new paragraphs after innovation table")

# ===== Step 4: Modify the abstract to emphasize innovation =====
# Find the abstract paragraph (element #007 in earlier inspection, contains "本项目核心算法")
# Add a new sentence about innovation features
abstract_modified = False
for el in elements:
    tag = el.tag.split('}')[-1]
    if tag == 'p':
        from docx.text.paragraph import Paragraph
        p = Paragraph(el, doc)
        text = p.text
        if '项目核心算法' in text and '快速傅里叶变换' in text:
            # Append innovation emphasis to this paragraph
            # Add a new run at the end
            run = p.add_run(
                " 项目在任务书要求之外自主实现 8 项创新功能，其中生存模式（血量条机制）、"
                "主题系统（情绪自动分类）、谱面密度归一化（时长自适应）、难度自动评级（四维加权评分）"
                "为最具代表性的核心创新亮点。"
            )
            run.bold = True
            abstract_modified = True
            print(f"Modified abstract paragraph to emphasize innovation")
            break

if not abstract_modified:
    print("WARN: Abstract paragraph not found")

# ===== Step 5: Save =====
doc.save(DST)
print(f"\nSaved to: {DST}")
import os
print(f"File size: {os.path.getsize(DST)} bytes")
